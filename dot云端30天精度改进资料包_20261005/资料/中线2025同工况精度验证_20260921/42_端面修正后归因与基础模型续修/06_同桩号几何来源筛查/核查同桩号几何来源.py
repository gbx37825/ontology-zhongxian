"""只读原断面与原表同桩号几何覆盖；匹配不等于物理归属验收，未匹配不自动纠错。"""
from pathlib import Path
import collections
import csv
import hashlib
import json
import math
import openpyxl

HERE=Path(__file__).resolve().parent
R=HERE.parents[1]
SOURCE=R/'01_原始资料/中线全线建模数据.xlsx'
MODEL=R/'26_冻结基线与后续归因/共享输入/C07.json'
TOL=1e-6
sha=lambda p:hashlib.sha256(Path(p).read_bytes()).hexdigest()


def number(v):
    try:
        x=float(v)
        return x if math.isfinite(x) else None
    except (TypeError,ValueError):
        return None


def main():
    sources={str(p):sha(p) for p in [Path(__file__),SOURCE,MODEL]}
    wb=openpyxl.load_workbook(SOURCE,data_only=True,read_only=True)
    table=collections.defaultdict(list)
    for sheet in ['渠道','渐变段','倒虹吸','节制闸']:
        for rownum,row in enumerate(wb[sheet].iter_rows(values_only=True),1):
            if rownum==1:continue
            row=list(row)+[None]*12
            if sheet=='渠道':
                sides=[('in',2,6,4,5,8),('out',3,7,4,5,9)]
            elif sheet in ['渐变段','倒虹吸']:
                sides=[('in',2,3,4,5,None),('out',6,7,8,9,None)]
            else:
                sides=[('gate',2,6,None,None,None)]
            for side,ix,iz,ib,im,itop in sides:
                x,z=number(row[ix]),number(row[iz])
                if x is None or z is None:continue
                if sheet=='节制闸':
                    n,b=number(row[4]),number(row[5]);b=n*b if n is not None and b is not None else None;m=0.
                else:b,m=number(row[ib]),number(row[im])
                if b is None or m is None:continue
                top=number(row[itop]) if itop is not None else None
                position=x*1000
                table[round(position,6)].append(dict(sheet=sheet,row=rownum,name=row[1],side=side,
                    chainage_m=position,bed_m=z,bottom_width_m=b,side_slope=m,top_m=top,
                    source_geometry_cells=[openpyxl.utils.get_column_letter(i+1)+str(rownum)
                          for i in [iz,*([4,5] if sheet=='节制闸' else [ib,im])]],
                    qualification='闸表宽度=孔数×单孔宽，仅作旁证，不自动等同明渠横断面' if sheet=='节制闸' else '原表对应端点'))
    wb.close()
    src=json.loads(MODEL.read_text())['中线全线'];records=[]
    owners=collections.defaultdict(list)
    for name,c in src['baseData']['channels'].items():
        for sec in c['sec_names']:owners[sec].append(name)
    for name,sec in src['baseData']['sections'].items():
        pts=sec['yz'];z=min(p[1] for p in pts);top=max(p[1] for p in pts);height=top-z
        floor=[p[0] for p in pts if abs(p[1]-z)<=TOL]
        if len(floor)<2 or height<=0:
            records.append(dict(section=name,status='UNSUPPORTED_DIRECT_PROFILE_EXTRACTION'));continue
        b=max(floor)-min(floor);m=(max(p[0] for p in pts)-min(p[0] for p in pts)-b)/(2*height)
        candidates=table.get(round(float(sec['location']),6),[])
        matches=[c for c in candidates if abs(c['bed_m']-z)<=TOL and abs(c['bottom_width_m']-b)<=TOL and abs(c['side_slope']-m)<=TOL]
        channel_matches=[c for c in matches if c['sheet']!='节制闸']
        tops=[c for c in channel_matches if c['top_m'] is not None]
        status='SOURCE_GEOMETRY_COVERED' if channel_matches else ('GATE_ONLY_GEOMETRY_SUPPORT' if matches else ('NO_SOURCE_CHAINAGE' if not candidates else 'NO_SAME_CHAINAGE_GEOMETRY_MATCH'))
        records.append(dict(section=name,owners=owners[name],chainage_m=sec['location'],bed_m=z,
             bottom_width_m=b,side_slope=m,top_m=top,status=status,source_candidates=candidates,
             matching_source_rows=matches,channel_top_candidates=tops,
             source_channel_top_match=any(abs(c['top_m']-top)<=TOL for c in tops) if tops else None))
    assert all(sha(p)==v for p,v in sources.items())
    flagged=[x for x in records if x['status']!='SOURCE_GEOMETRY_COVERED' or x.get('source_channel_top_match') is False]
    summary=dict(scope='原1394断面同桩号候选覆盖；原表角色冲突和重复值不自动择优，非断面现场验收',
        records_count=len(records),status_counts=dict(collections.Counter(x['status'] for x in records)),
        matching_channel_top_checks=sum(x.get('source_channel_top_match') is not None for x in records),
        matching_channel_top_conflicts=sum(x.get('source_channel_top_match') is False for x in records),
        flagged_count=len(flagged),flagged=flagged,source_sha256=sources,tolerance_m=TOL,
        original_files_changed=False,hydraulic_steps=0,
        limitations=['只使用原表端点，不能凭最近桩号/多数相同值自动改原输入。',
                     '四点底宽/平均边坡提取仅用于已知原矩形梯形；命中任一同桩号候选不证明渠道归属正确。',
                     '源表堤顶仅在同位置且几何匹配的渠道行中检查；不能把闸顶、报警水位或邻点堤顶替换成断面顶部。'])
    (HERE/'同桩号几何来源覆盖.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2)+'\n')
    (HERE/'全部原断面来源对照.json').write_text(json.dumps(records,ensure_ascii=False,indent=2)+'\n')
    with (HERE/'待核原断面简表.csv').open('w') as f:
        fields=['section','chainage_m','bed_m','bottom_width_m','side_slope','top_m','status','source_channel_top_match']
        writer=csv.DictWriter(f,fieldnames=fields,extrasaction='ignore');writer.writeheader();writer.writerows(flagged)
    print(json.dumps({k:v for k,v in summary.items() if k not in ['source_sha256','flagged','limitations']},ensure_ascii=False))


if __name__=='__main__':main()
