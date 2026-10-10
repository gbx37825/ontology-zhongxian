from pathlib import Path
import collections,hashlib,json,zipfile
import openpyxl

O=Path(__file__).resolve().parent;R=O.parents[1]
source=Path('/Users/gaobingxiang/code/hydroclaw0418/datasources/中线')
workbooks=[source/n for n in ['中线全线建模数据.xlsx','中线建模编号说明.xlsx','中线组件信息.xlsx']]+[R/'01_原始资料'/n for n in ['流量计退水闸节制闸位置统计.xlsx','附件：中线各闸站水位计和流量计安装位置排查表 (补充).xlsx']]
keywords=['局部水力','局部损失','速度水头','水准点','校准','标定','1985','黄海','大沽','竣工','仪器编号','正向','反向','收缩系数','扩散系数']
books=[]
for p in workbooks:
    w=openpyxl.load_workbook(p,read_only=True,data_only=True);hits=[]
    for s in w:
        for row in s:
            for c in row:
                if isinstance(c.value,str) and any(k in c.value for k in keywords):hits.append(dict(sheet=s.title,cell=c.coordinate,text=c.value))
    books.append(dict(path=str(p),sha256=hashlib.sha256(p.read_bytes()).hexdigest(),sheet_dimensions={s.title:[s.max_row,s.max_column]for s in w},keyword_hits=hits));w.close()
archives=[]
for p in list(source.glob('*.zip'))+list((source/'2025年数据').glob('*.zip')):
    with zipfile.ZipFile(p)as z:
        names=[v for v in z.namelist()if not v.endswith('/')]
        archives.append(dict(path=str(p),entries=len(names),types=dict(collections.Counter(Path(v).suffix.lower()for v in names)),drawing_or_datum_named_entries=[v for v in names if any(k in v for k in ['竣工','纵剖','横断','水准','基准','校准'])]))
inventory=[]
for parent in [source,R/'01_原始资料']:
    for p in parent.iterdir():
        if p.is_file():inventory.append(dict(path=str(p),bytes=p.stat().st_size))
documents=[R/'22_马峰论文方法核对/论文段落摘录.txt',R/'29_长距离明渠方案方法核对/方案逐页摘录.txt']
docevidence=[]
for p in documents:
    lines=p.read_text().splitlines()
    hits=[dict(line=i,text=line)for i,line in enumerate(lines,1)if any(k in line for k in ['局部损失系数','速度水头','竣工','水准点','CH152','CH320'])]
    docevidence.append(dict(path=str(p),hits=hits,sha256=hashlib.sha256(p.read_bytes()).hexdigest()))
result=dict(scope='限定用户中线datasources、已有01原始副本、已摘录两份PDF/论文；未扫整机。压缩包只检查目录清单，不据此保证每张运行截图无任何附注。',
    workbooks=books,archives=archives,direct_files=inventory,document_evidence=docevidence,
    coefficient_definition_found=False,facility_specific_velocity_reference_found=False,
    required_materials=[
      '渐变段K/L两列原数据字典或原求解器读取段及实际水头损失公式：单位、触发流态/流向、参考流速断面、是否两列同时使用。',
      '倒虹吸源loss_coeff=0.1的适用范围说明或率定记录：是否已含进口/出口渐变、闸门或弯头。',
      'CH152 280+528至285+398，含北汝河出口渐变、CH152及青龙河进出口的竣工纵剖面和横断面；明确渠顶/墙顶/一级马道、盖板和边坡。',
      'CH320 580+999至585+487，含小凹沟出口渐变、CH320、石门河入口渐变的竣工纵剖面和横断面；核实89.241/89.202m渠底及98.341/98.302m顶高程。',
      '新蟒河485+755至485+850、玉带河266+405至266+475进口前池/渐变段的墙顶、顶板、开敞有压分界图及高程基准。',
      'JZZ15 K280+683左右水位计及其他13个超轮廓测点的2025年仪器序列号、历史列映射、测点迁移记录、水准点引测与零点复核记录、统一高程基准换算。',
      '西赵河69+521至69+872及槐河(二)923km附近的竣工统一桩号对照；中易水进口渐变段1145+907至1146+974的真实长度及节点概化说明。',
    ])
(O/'资料范围与缺项.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
print(json.dumps({'books':len(books),'archives':archives,'definition_found':False},ensure_ascii=False,indent=2))
